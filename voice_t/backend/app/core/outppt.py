import os
import base64
from pptx import Presentation
from openai import OpenAI
import json
import sys
from app.core.app_prompts import PPT_EXTRACTION_PROMPTS

class PPTExtractor:
    def __init__(self, api_key=None, base_url=None, model="qwen-omni-turbo"):
        """
        初始化PPT提取器

        Args:
            api_key: API密钥
            base_url: API基础URL (例如DashScope的URL)
            model: 使用的模型名称
        """
        # 从环境变量或参数中获取API密钥
        self.api_key = api_key or os.environ.get("OPENAI_API_KEY", "sk-9806a9e11e6a4a4daf6e588bbed24fcf")
        self.base_url = base_url or os.environ.get("OPENAI_BASE_URL", "https://dashscope.aliyuncs.com/compatible-mode/v1")
        self.model = model

        # 初始化OpenAI客户端
        self.client = OpenAI(
            api_key=self.api_key,
            base_url=self.base_url,
        )

        # 获取提示词配置
        self.system_prompt = PPT_EXTRACTION_PROMPTS["system_prompt"]
        self.image_extraction_prompt = PPT_EXTRACTION_PROMPTS["image_extraction_prompt"]
        self.slide_summary_prompt = PPT_EXTRACTION_PROMPTS["slide_summary_prompt"]

        # 创建用于保存提取结果的字典
        self.extracted_content = {
            "metadata": {
                "total_slides": 0,
                "has_images": False,
                "extraction_time": None
            },
            "slides": []
        }
    def get_slide_count(self,pptx_path):
        """获取PPT的幻灯片总数"""
        prs = Presentation(pptx_path)
        return len(prs.slides)
    def encode_image(self, image_path):
        """将图片转换为Base64编码"""
        with open(image_path, "rb") as image_file:
            return base64.b64encode(image_file.read()).decode("utf-8")

    def extract_text_from_pptx(self, pptx_path, output_format="txt", save_images=True):
        """
        从PPTX文件中提取文本和图片内容

        Args:
            pptx_path: PPTX文件路径
            output_format: 输出格式 ("txt" 或 "json")
            save_images: 是否保存图片文件

        Returns:
            提取的内容 (字符串或字典)
        """
        import datetime

        prs = Presentation(pptx_path)
        extracted_text = []
        img_folder = "ppt_images"

        if save_images:
            os.makedirs(img_folder, exist_ok=True)

        # 更新元数据
        self.extracted_content["metadata"]["total_slides"] = len(prs.slides)
        self.extracted_content["metadata"]["extraction_time"] = datetime.datetime.now().isoformat()

        for slide_idx, slide in enumerate(prs.slides):
            slide_num = slide_idx + 1
            slide_content = {
                "slide_number": slide_num,
                "text_content": "",
                "images": []
            }

            slide_text = f"\n--- 幻灯片 {slide_num} ---\n"

            # 1. 提取幻灯片文本
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text.strip():
                    slide_text += shape.text + "\n"
                    slide_content["text_content"] += shape.text + "\n"

            extracted_text.append(slide_text)

            # 2. 处理幻灯片中的图片
            image_found = False
            for shape_idx, shape in enumerate(slide.shapes):
                if hasattr(shape, "image"):
                    image_found = True
                    self.extracted_content["metadata"]["has_images"] = True

                    img_path = os.path.join(img_folder, f"slide_{slide_num}_img_{shape_idx + 1}.png")
                    img_description = ""

                    # 保存图片
                    if save_images:
                        with open(img_path, "wb") as f:
                            f.write(shape.image.blob)

                    # 通过API解析图片内容
                    base64_image = self.encode_image(img_path) if save_images else base64.b64encode(shape.image.blob).decode("utf-8")

                    completion = self.client.chat.completions.create(
                        model=self.model,
                        messages=[
                            {"role": "system", "content": [{"type": "text", "text": self.system_prompt}]},
                            {"role": "user", "content": [
                                {"type": "image_url", "image_url": {"url": f"data:image/png;base64,{base64_image}"}},
                                {"type": "text", "text": self.image_extraction_prompt},
                            ]},
                        ],
                        modalities=["text"],
                        stream=True,
                        stream_options={"include_usage": True},
                    )

                    full_content = ""
                    for chunk in completion:
                        if chunk.choices:
                            delta = chunk.choices[0].delta
                            if delta.content:
                                full_content += delta.content

                    img_description = full_content.strip()
                    if img_description:
                        image_text = f"【图片内容 - 幻灯片 {slide_num}】: {img_description}\n"
                        extracted_text.append(image_text)

                        # 添加到JSON结构
                        slide_content["images"].append({
                            "image_index": shape_idx + 1,
                            "image_path": img_path if save_images else None,
                            "description": img_description
                        })

            # 将幻灯片内容添加到结果中
            self.extracted_content["slides"].append(slide_content)

        # 根据输出格式返回结果
        if output_format == "json":
            return self.extracted_content
        else:
            return "\n".join(extracted_text)

    def save_output(self, content, output_path, format="txt"):
        """保存提取的内容到文件"""
        if format == "json":
            with open(output_path, "w", encoding="utf-8") as f:
                json.dump(content, f, ensure_ascii=False, indent=2)
        else:
            with open(output_path, "w", encoding="utf-8") as f:
                f.write(content)

        # 内容已保存

    def run(self, pptx_path, output_path=None, output_format="txt", save_images=True):
        """运行PPT内容提取程序

        Args:
            pptx_path: PPTX文件路径
            output_path: 输出文件路径
            output_format: 输出格式 ("txt" 或 "json")
            save_images: 是否保存图片文件

        Returns:
            提取的内容 (字符串或字典), 幻灯片总数
        """
        slide_count = self.get_slide_count(pptx_path)
        # 如果未提供输出路径，则生成默认路径
        if not output_path:
            base_name = os.path.splitext(os.path.basename(pptx_path))[0]
            output_path = f"{base_name}_extracted.{output_format}"

        # 提取内容
        content = self.extract_text_from_pptx(
            pptx_path,
            output_format=output_format,
            save_images=save_images
        )

        # 保存内容
        if output_path:
            self.save_output(content, output_path, format=output_format)

        return content, slide_count

# 当作为主程序运行时
if __name__ == "__main__":
    import argparse

    parser = argparse.ArgumentParser(description="从PPT中提取文本和图片内容")
    parser.add_argument("pptx_file", help="PPTX文件路径")
    parser.add_argument("--output", "-o", help="输出文件路径")
    parser.add_argument("--format", "-f", choices=["txt", "json"], default="txt", help="输出格式 (txt 或 json)")
    parser.add_argument("--no-images", action="store_true", help="不保存图片文件")
    parser.add_argument("--api-key", help="API密钥")
    parser.add_argument("--api-url", help="API基础URL")
    parser.add_argument("--model", default="qwen-omni-turbo", help="使用的模型名称")

    args = parser.parse_args()

    extractor = PPTExtractor(
        api_key=args.api_key,
        base_url=args.api_url,
        model=args.model
    )

    content,slide_count = extractor.run(
        args.pptx_file,
        output_path=args.output,
        output_format=args.format,
        save_images=not args.no_images
    )
    print(f"PPT总页数: {slide_count}")